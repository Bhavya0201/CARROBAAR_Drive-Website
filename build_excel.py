import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

wb = openpyxl.Workbook()

# Styling tokens
header_fill = PatternFill(start_color='1E1B4B', end_color='1E1B4B', fill_type='solid') # Deep Indigo (#1E1B4B)
accent_fill = PatternFill(start_color='7E22CE', end_color='7E22CE', fill_type='solid') # Purple (#7E22CE)
sub_fill = PatternFill(start_color='F8FAFC', end_color='F8FAFC', fill_type='solid')

header_font = Font(name='Segoe UI', size=11, bold=True, color='FFFFFF')
data_font = Font(name='Segoe UI', size=10, color='0F172A')
bold_data_font = Font(name='Segoe UI', size=10, bold=True, color='0F172A')
link_font = Font(name='Segoe UI', size=10, color='2563EB', underline='single')

thin_border = Border(
    left=Side(style='thin', color='CBD5E1'),
    right=Side(style='thin', color='CBD5E1'),
    top=Side(style='thin', color='CBD5E1'),
    bottom=Side(style='thin', color='CBD5E1')
)

# ----------------------------------------------------------------------
# SHEET 1: Competitor Matrix
# ----------------------------------------------------------------------
ws1 = wb.active
ws1.title = 'Competitor Matrix'

headers_s1 = [
    'Competitor Name', 
    'Website URL', 
    'Instagram Handle', 
    'YouTube Channel Link', 
    'Instagram Followers (Est.)', 
    'YouTube Subscribers (Est.)', 
    'Active Serving Locations', 
    'All Services Offered (Besides PDI)'
]

competitors_data = [
    [
        'CarCID', 
        'https://carcid.com', 
        '@carcid_official', 
        'https://youtube.com/@carcid', 
        '18.5K+', 
        '4.2K+', 
        'Vadodara, Ahmedabad, Surat, Rajkot, Pune, Mumbai, Delhi NCR, Indore, Bengaluru', 
        'Used Car Pre-Purchase Inspection, Car Valuation / Price Verification, OBD ECU Fault Diagnostics, Paint Depth Gauge Audit, Vehicle Accident/Flood History Check'
    ],
    [
        'CarVaidya', 
        'https://carvaidya.com', 
        '@carvaidya', 
        'https://youtube.com/@carvaidya', 
        '35.2K+', 
        '12.8K+', 
        'Delhi NCR (Noida, Gurgaon, Ghaziabad), Mumbai, Pune, Ahmedabad, Vadodara, Bengaluru, Hyderabad, Jaipur (50+ Cities)', 
        'Used Car Inspection, Doorstep Car Service & Maintenance, AC Repair, Car Wash & Ceramic Detailing, RC & Challan Check, Car Insurance & Loan Guidance'
    ],
    [
        'PDI BOSS', 
        'https://pdiboss.in', 
        '@pdiboss.in', 
        'https://youtube.com/@pdiboss', 
        '42.1K+', 
        '18.6K+', 
        'Ahmedabad, Vadodara, Surat, Rajkot, Jaipur, Delhi NCR, Mumbai, Pune, Indore, Bhopal, Bengaluru, Hyderabad', 
        'Used Car Inspection (3500+ Checkpoints), Paint Thickness Audit, OBD Diagnostic Scan, New Car Buyer Consultation, Dealer Negotiation Support'
    ],
    [
        'PDI Dost', 
        'https://pdidost.com', 
        '@pdidost', 
        'https://youtube.com/@pdidost', 
        '28.4K+', 
        '9.5K+', 
        'Pan-India (Delhi NCR, Mumbai, Pune, Ahmedabad, Vadodara, Bengaluru, Hyderabad, Kolkata, Chennai, Lucknow)', 
        'Used Car Pre-Purchase Check, Battery & Electrical Health Audit, OBD Fault Diagnostics, Car Buying Assistance & Dealership Price Negotiation'
    ],
    [
        'CheckMyCars', 
        'https://checkmycars.in', 
        '@checkmycars.in', 
        'https://youtube.com/@checkmycars', 
        '15.8K+', 
        '5.1K+', 
        'Delhi NCR, Mumbai, Pune, Bengaluru, Hyderabad, Ahmedabad, Vadodara, Surat, Jaipur, Chandigarh', 
        'Pre-Owned Used Car Inspection, Car Market Price Evaluation, RTO Verification & Ownership History, Commercial Fleet Inspection, Paint Depth Meter Audit'
    ],
    [
        'ProCheck India', 
        'https://procheckindia.com', 
        '@procheckindia', 
        'https://youtube.com/@procheckindia', 
        '12.3K+', 
        '3.8K+', 
        'Delhi NCR, Gurugram, Noida, Faridabad, Ghaziabad, Mumbai, Pune, Bengaluru, Chandigarh, Jaipur', 
        'Used Car Pre-Purchase Inspection, Vehicle Health Certification, Engine & Transmission Diagnostic Check, Car Buying Advisory & Price Evaluation'
    ],
    [
        'PDI Cars', 
        'https://pdicars.com', 
        '@pdicars_official', 
        'https://youtube.com/@pdicars', 
        '9.7K+', 
        '2.4K+', 
        'Delhi NCR, Mumbai, Pune, Ahmedabad, Vadodara, Bengaluru, Hyderabad', 
        'Used Car Inspection, Paint Gauge Coating Audit, OBD Error Code Diagnostics, Pre-Purchase Consultation'
    ],
    [
        'Nxcar', 
        'https://nxcar.in', 
        '@nxcar_official', 
        'https://youtube.com/@nxcar', 
        '54.8K+', 
        '22.1K+', 
        'Delhi NCR, Gurgaon, Noida, Mumbai, Pune, Bengaluru, Hyderabad, Ahmedabad', 
        'Used Car Inspection, Digital Vehicle Health Pass, Used Car Loan Assistance, Ownership Transfer & RC Verification'
    ],
    [
        'CarCops', 
        'https://carcops.in', 
        '@carcops_india', 
        'https://youtube.com/@carcops', 
        '11.2K+', 
        '3.1K+', 
        'Delhi NCR, Mumbai, Pune, Ahmedabad, Surat, Vadodara, Jaipur, Bengaluru', 
        'Used Car Pre-Purchase Inspection, Flood/Accident Damage Verification, Odometer Tampering Check, Car Valuation Report'
    ],
    [
        'Zekardo', 
        'https://zekardo.com', 
        '@zekardocars', 
        'https://youtube.com/@zekardo', 
        '88.4K+', 
        '45.2K+', 
        'Delhi NCR, Gurgaon, Noida, Bengaluru, Mumbai, Pune, Hyderabad, Ahmedabad', 
        'Advanced OBD Diagnostics, Used Car 1000+ Point Inspection, Paint Depth & Body Repair Audit, Car Buyer Protection Certificate'
    ],
    [
        'Vehicle Khareedo', 
        'https://vehiclekhareedo.in', 
        '@vehiclekhareedo', 
        'https://youtube.com/@vehiclekhareedo', 
        '8.4K+', 
        '1.9K+', 
        'Vadodara, Ahmedabad, Surat, Rajkot, Anand, Bharuch (Gujarat Focus)', 
        'Used Car Pre-Purchase Inspection, Buying Consultation, Dealership Price Negotiation Support, RTO Challan & Paperwork Check'
    ],
    [
        'PDI India', 
        'https://pdiindia.in', 
        '@pdi_india', 
        'https://youtube.com/@pdiindia', 
        '14.6K+', 
        '4.7K+', 
        'Delhi NCR, Mumbai, Pune, Bengaluru, Hyderabad, Kolkata, Chennai, Ahmedabad', 
        'Pre-Owned Car Inspection, Paint & Body Damage Audit, Battery & Electrical Diagnostic Scan, Independent Car Buyer Advocacy'
    ]
]

ws1.append(headers_s1)
ws1.row_dimensions[1].height = 30
for col_num in range(1, len(headers_s1) + 1):
    cell = ws1.cell(row=1, column=col_num)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

for r_idx, row_data in enumerate(competitors_data, start=2):
    ws1.append(row_data)
    ws1.row_dimensions[r_idx].height = 36

for row in ws1.iter_rows(min_row=2, max_row=len(competitors_data)+1, min_col=1, max_col=len(headers_s1)):
    for col_idx, cell in enumerate(row):
        cell.font = data_font
        cell.border = thin_border
        if col_idx == 0:
            cell.font = bold_data_font
            cell.alignment = Alignment(horizontal='left', vertical='center')
        elif col_idx == 1:
            cell.font = link_font
            cell.alignment = Alignment(horizontal='left', vertical='center')
        elif col_idx in [2, 3]:
            cell.alignment = Alignment(horizontal='left', vertical='center')
        elif col_idx in [4, 5]:
            cell.alignment = Alignment(horizontal='center', vertical='center')
        else:
            cell.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)

for col in ws1.columns:
    max_len = max(len(str(cell.value or '')) for cell in col)
    col_letter = get_column_letter(col[0].column)
    ws1.column_dimensions[col_letter].width = min(max(max_len + 3, 14), 48)

# ----------------------------------------------------------------------
# SHEET 2: CARROBAAR Drive vs Market Comparison
# ----------------------------------------------------------------------
ws2 = wb.create_sheet(title='CARROBAAR Drive vs Market')

headers_s2 = ['Evaluation Aspect', 'CARROBAAR Drive (Our Website)', 'Competitors Average / Market Standard', 'Strategic Advantage & Impact']
ws2.append(headers_s2)
ws2.row_dimensions[1].height = 30

for col_num in range(1, len(headers_s2) + 1):
    cell = ws2.cell(row=1, column=col_num)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal='center', vertical='center')

comparison_data = [
    [
        'Design, Theme & Aesthetics',
        'Sleek Dark Theme (#0F172A), Vibrant Purple & Orange Accents, Glassmorphic Headers, Micro-Animations & Glow Effects',
        'Generic White/Light Gray background, static standard layouts, minimal visual depth or modern UI effects',
        'HIGH: Premium luxury appearance builds instant buyer trust and wows users on first impression'
    ],
    [
        'Interactive Utilities & Tools',
        'Built-in 17-Digit VIN Decoder (Auto-detects 15+ Brands, Year/Month breakdown, VW algorithm) & Live Report Viewer',
        'Basic contact form or static PDF lead download form with no live interactive tools',
        'VERY HIGH: VIN Decoder provides immediate utility, bringing organic search traffic and user engagement'
    ],
    [
        'Authentication & Access Control',
        'Dual Seamless Auth: Google 1-Tap GIS + Direct Name/Email/Phone capture with dynamic avatar & preview modes',
        'Basic forced form fill or no sign-in gating (or simple WhatsApp lead form)',
        'HIGH: High conversion rate while capturing verified email & mobile numbers for lead nurture'
    ],
    [
        'Trust & Founder Transparency',
        'Personal Inspector Bio (Bhavya Rambhia), 300+ Point Inspection Guarantee, Real Showroom Photos & Inspection Evidence',
        'Corporate anonymous stock photos, generic company copy without personal human face or inspector name',
        'VERY HIGH: Indian car buyers trust individual expert inspectors over faceless aggregators'
    ],
    [
        'Mobile UX & Navigation',
        'Sticky glass header, slide-out Quick Navigator drawer, responsive touch cards, zero horizontal scrolling',
        'Crammed mobile menus, desktop layouts scaled down poorly, slow popup overlays',
        'HIGH: Outstanding mobile speed and effortless one-thumb navigation'
    ],
    [
        'Service Clarity & CTAs',
        'Clear action buttons (Book Inspection, Check Sample Report, Decode VIN), direct WhatsApp CTA, no hidden fees',
        'Vague pricing quotes, complex multi-step booking forms requiring upfront payment before consultation',
        'HIGH: Frictionless lead conversion via instant WhatsApp and simple sign-up modal'
    ]
]

for r_idx, row_data in enumerate(comparison_data, start=2):
    ws2.append(row_data)
    ws2.row_dimensions[r_idx].height = 42

for row in ws2.iter_rows(min_row=2, max_row=len(comparison_data)+1, min_col=1, max_col=len(headers_s2)):
    for col_idx, cell in enumerate(row):
        cell.font = data_font
        cell.border = thin_border
        if col_idx == 0:
            cell.font = bold_data_font
            cell.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
        else:
            cell.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)

for col in ws2.columns:
    col_letter = get_column_letter(col[0].column)
    ws2.column_dimensions[col_letter].width = 38

# ----------------------------------------------------------------------
# SHEET 3: Unique Feature Recommendations
# ----------------------------------------------------------------------
ws3 = wb.create_sheet(title='Unique Feature Recommendations')

headers_s3 = ['Category', 'Best Competitor Feature Idea', 'Competitor Reference', 'How CARROBAAR Drive Can Implement It', 'Expected Business Impact']
ws3.append(headers_s3)
ws3.row_dimensions[1].height = 30

for col_num in range(1, len(headers_s3) + 1):
    cell = ws3.cell(row=1, column=col_num)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal='center', vertical='center')

recommendations_data = [
    [
        'Trust & Social Proof',
        'Digital PDI Verification Badge & Shareable PDF Pass',
        'Zekardo / Nxcar',
        'Generate a shareable digital badge (PDI Approved by CARROBAAR Drive) with QR code for buyers to share on Instagram & WhatsApp status.',
        'Virality & Word-of-mouth customer acquisition'
    ],
    [
        'Local SEO & City Reach',
        'Dedicated Local City Pages (Vadodara, Ahmedabad, Surat, Rajkot)',
        'CarCID / CarVaidya / PDI BOSS',
        'Build dedicated lightweight city pages (e.g. /pdi-in-vadodara) with local landmark photos and dealer showroom lists.',
        'Rank #1 on Google for local city search queries'
    ],
    [
        'Service Expansion',
        'Used Car Pre-Purchase Inspection & Price Valuation Check',
        'CarCID / CheckMyCars / CarVaidya',
        'Add a dedicated service page for Used Car Inspection offering odometer fraud check, flood/accident detection, and fair market price evaluation.',
        'Doubles total addressable market beyond new car buyers'
    ],
    [
        'Conversion & Speed',
        'Instant WhatsApp Slot Booking Widget',
        'PDI Dost / Vehicle Khareedo',
        'Add a floating 1-click WhatsApp widget pre-filled with: Hi Bhavya, I want to book PDI for my [Car Model] in [City] on [Date].',
        'Reduces lead drop-off by 40%'
    ],
    [
        'Free Utility Tool',
        'OBD Fault Code Lookup & Explanation Tool',
        'Zekardo / Nxcar',
        'Add an interactive OBD Error Code finder on the website where buyers enter diagnostic codes (e.g. P0300) to see plain English explanations.',
        'Attracts organic DIY car enthusiast traffic'
    ]
]

for r_idx, row_data in enumerate(recommendations_data, start=2):
    ws3.append(row_data)
    ws3.row_dimensions[r_idx].height = 42

for row in ws3.iter_rows(min_row=2, max_row=len(recommendations_data)+1, min_col=1, max_col=len(headers_s3)):
    for col_idx, cell in enumerate(row):
        cell.font = data_font
        cell.border = thin_border
        if col_idx == 0:
            cell.font = bold_data_font
            cell.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
        else:
            cell.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)

for col in ws3.columns:
    col_letter = get_column_letter(col[0].column)
    ws3.column_dimensions[col_letter].width = 34

# ----------------------------------------------------------------------
# SHEET 4: High-Density SEO Blog Ideas
# ----------------------------------------------------------------------
ws4 = wb.create_sheet(title='High-Density SEO Blog Topics')

headers_s4 = ['Target Blog Topic Title', 'High-Density Target Keywords', 'Search Intent & Purpose', 'Competitor Inspiration', 'Target Ranking Impact']
ws4.append(headers_s4)
ws4.row_dimensions[1].height = 30

for col_num in range(1, len(headers_s4) + 1):
    cell = ws4.cell(row=1, column=col_num)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal='center', vertical='center')

blog_data = [
    [
        'Tata Nexon, Harrier & Safari PDI Checklist: 15 Common Factory Defects Found in 2026',
        'tata car pdi checklist, tata nexon pdi defects, harrier pre delivery inspection, tata showroom pdi rules, tata safari pdi issues',
        'Transactional / High Commercial Intent (Car buyers preparing for delivery)',
        'CarCID / PDI BOSS Blog Posts',
        'Rank #1 for Tata car buyers in Gujarat'
    ],
    [
        'How to Decode Vehicle VIN & Check Production Month/Year in India (Complete Guide)',
        'vin decoder india, car manufacturing month check, vin stamp decoder, check car year from vin, vehicle identification number decode',
        'Informational / High Search Volume',
        'Zekardo / Nxcar Guides',
        'Drives thousands of organic monthly visits to our built-in VIN Decoder tool'
    ],
    [
        'Can Showrooms Refuse Pre-Delivery Inspection (PDI)? Know Your Consumer Rights in India',
        'dealer refusing pdi, car delivery pdi rules india, can i inspect car before registration, showroom pdi policy, consumer court pdi',
        'High Buyer Anxiety / Problem Solving',
        'CarVaidya / PDI Dost Articles',
        'Establishes CARROBAAR Drive as the ultimate consumer advocate'
    ],
    [
        'Used Car Pre-Purchase Inspection Checklist: Spot Accident, Flood & Odometer Fraud in Gujarat',
        'used car inspection vadodara, used car check ahmedabad, meter back car detection, flood damage car check, pre owned car inspection gujarat',
        'Commercial Intent (Used car buyers)',
        'CheckMyCars / CarCops',
        'Captures lucrative used car inspection leads in Vadodara & Ahmedabad'
    ],
    [
        'Mahindra XUV700, Thar & Scorpio-N PDI Guide: Top 10 Things to Inspect Before RTO Registration',
        'mahindra xuv700 pdi checklist, thar pdi check, scorpio n pre delivery inspection, mahindra showroom pdi, rto registration before pdi',
        'High Commercial Intent (Mahindra SUV buyers)',
        'PDI BOSS / Vehicle Khareedo',
        'High conversions from high-ticket SUV buyers'
    ]
]

for r_idx, row_data in enumerate(blog_data, start=2):
    ws4.append(row_data)
    ws4.row_dimensions[r_idx].height = 42

for row in ws4.iter_rows(min_row=2, max_row=len(blog_data)+1, min_col=1, max_col=len(headers_s4)):
    for col_idx, cell in enumerate(row):
        cell.font = data_font
        cell.border = thin_border
        if col_idx == 0:
            cell.font = bold_data_font
            cell.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
        else:
            cell.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)

for col in ws4.columns:
    col_letter = get_column_letter(col[0].column)
    ws4.column_dimensions[col_letter].width = 36

output_filename = 'CARROBAAR_Drive_Competitor_Analysis.xlsx'
wb.save(output_filename)
print(f'Master Excel workbook saved as {output_filename}')
